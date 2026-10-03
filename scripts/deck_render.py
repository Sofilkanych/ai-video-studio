"""PPTX з плану слайдів (режим «текст → презентація»): теми, 8 типів слайдів, нотатки доповідача.

Тема «forest» відтворює стиль із реальної лекції: Georgia + Calibri, темно-зелений і теплий
світлий фон. Розміри тексту підбираються під рамку (python-pptx не вміє autofit без шрифтів).
"""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

THEMES: dict[str, dict[str, Any]] = {
    "forest": {
        "heading_font": "Georgia", "body_font": "Calibri",
        "dark_bg": "1B2622", "light_bg": "F4F3F0",
        "ink": "1B2622", "body": "3E4A45", "muted": "77827D",
        "on_dark": "FFFFFF", "on_dark_soft": "CFE3DC", "on_dark_muted": "A9B5B0",
        "accent": "1D6F5C", "card": "DCDDD8", "card_soft": "E7F0ED", "card_dark": "27332E",
    },
    "neutral": {
        "heading_font": "Georgia", "body_font": "Calibri",
        "dark_bg": "1F2933", "light_bg": "FFFFFF",
        "ink": "111827", "body": "374151", "muted": "6B7280",
        "on_dark": "FFFFFF", "on_dark_soft": "D1D5DB", "on_dark_muted": "9CA3AF",
        "accent": "2563EB", "card": "E5E7EB", "card_soft": "EFF6FF", "card_dark": "374151",
    },
}

W, H = 10.0, 5.625          # дюйми, 16:9
MX = 0.6                    # бокове поле
CHAR_W = 0.5                # середня ширина символу в частках кегля: Calibri, кирилиця
HEAD_CHAR_W = 0.57          # Georgia Bold ширша — інакше довгі заголовки переносяться на рядок більше
LINE_H = 1.22               # міжрядковий інтервал


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def text_height(text: str, width_in: float, size: float, char_w: float = CHAR_W) -> float:
    """Висота тексту в дюймах (оцінка за середньою шириною символу)."""
    chars_per_line = max(1, int(width_in * 72 / (size * char_w)))
    lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in text.split("\n"))
    return lines * size * LINE_H / 72


def fit_size(text: str, width_in: float, height_in: float, max_pt: float, min_pt: float = 9,
             char_w: float = CHAR_W) -> float:
    """Найбільший кегль, з яким текст уміщається в рамку (оцінка за середньою шириною символу)."""
    size = max_pt
    while size > min_pt:
        if text_height(text, width_in, size, char_w) <= height_in:
            return size
        size -= 0.5
    return min_pt





class Deck:
    def __init__(self, theme: str = "forest", footer: str = ""):
        self.t = THEMES[theme]
        self.footer = footer
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self._blank = self.prs.slide_layouts[6]

    # ---------- примітиви ----------

    def _slide(self, dark: bool):
        s = self.prs.slides.add_slide(self._blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(self.t["dark_bg"] if dark else self.t["light_bg"])
        return s

    def _box(self, s, x, y, w, h, fill: str, rounded: bool = False):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded:
            shape.adjustments[0] = 0.08
        return shape

    def _text(self, s, x, y, w, h, text: str, size: float, color: str, heading: bool = False,
              bold: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fit: bool = True,
              min_pt: float = 9, caps: bool = False):
        if caps:
            text = text.upper()
        if fit:
            size = fit_size(text, w, h, size, min_pt, HEAD_CHAR_W if heading else CHAR_W)
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = 1.1
            r = p.add_run()
            r.text = line
            f = r.font
            f.name = self.t["heading_font"] if heading else self.t["body_font"]
            f.size = Pt(size)
            f.bold = bold or heading
            f.color.rgb = _rgb(color)
        return tb

    def _badge(self, s, x, y, n: int | str, dark: bool = False):
        self._box(s, x, y, 0.42, 0.42, self.t["accent"])
        self._text(s, x, y, 0.42, 0.42, str(n), 15, "FFFFFF", bold=True, align=PP_ALIGN.CENTER,
                   anchor=MSO_ANCHOR.MIDDLE, fit=False)

    def _header(self, s, kicker: str, title: str, lead: str | None, dark: bool) -> float:
        """Надзаголовок, заголовок, підзаголовок; повертає y, з якого починається вміст."""
        t = self.t
        if kicker:
            self._text(s, MX, 0.34, W - 2 * MX, 0.24, kicker, 10, t["on_dark_soft"] if dark else t["accent"],
                       bold=True, caps=True, fit=False)
        self._text(s, MX, 0.62, W - 2 * MX, 0.52, title, 27, t["on_dark"] if dark else t["ink"],
                   heading=True, min_pt=18)
        y = 1.22
        if lead:
            self._text(s, MX, 1.18, W - 2 * MX, 0.42, lead, 13, t["on_dark_muted"] if dark else t["body"], min_pt=10)
            y = 1.72
        return y

    def _footer(self, s, n: int, dark: bool):
        c = self.t["on_dark_muted"] if dark else self.t["muted"]
        if self.footer:
            self._text(s, MX, 5.17, 5.0, 0.26, self.footer, 9, c, fit=False)
        self._text(s, 8.4, 5.17, 1.0, 0.26, str(n), 9, c, align=PP_ALIGN.RIGHT, fit=False)

    @staticmethod
    def _notes(s, narration: str):
        s.notes_slide.notes_text_frame.text = narration

    # ---------- типи слайдів ----------

    def title(self, d: dict[str, Any]):
        """Висота заголовка рахується з кількості рядків — підзаголовок і опис ідуть під ним."""
        s, t = self._slide(dark=True), self.t
        w, y = 6.6, 1.45
        if d.get("kicker"):
            self._text(s, MX, y - 0.37, w, 0.26, d["kicker"], 10.5, t["on_dark_soft"], bold=True, caps=True, fit=False)
        size = fit_size(d["title"], w, 1.5, 42, 26, HEAD_CHAR_W)
        h = text_height(d["title"], w, size, HEAD_CHAR_W)
        self._text(s, MX, y, w, h + 0.05, d["title"], size, t["on_dark"], heading=True, fit=False)
        y += h + 0.3
        if d.get("lead"):
            lsize = fit_size(d["lead"], w, 0.85, 22, 14, HEAD_CHAR_W)
            lh = text_height(d["lead"], w, lsize, HEAD_CHAR_W)
            self._text(s, MX, y, w, lh + 0.05, d["lead"], lsize, t["on_dark_soft"], heading=True, fit=False)
            y += lh + 0.3
        if d.get("meta"):
            self._text(s, MX, y, w, min(0.6, 4.95 - y), d["meta"], 13, t["on_dark_muted"], min_pt=10)
        return s

    def statement(self, d: dict[str, Any]):
        s, t = self._slide(dark=True), self.t
        self._header(s, d.get("kicker", ""), d["title"], None, dark=True)
        self._box(s, MX, 1.6, 0.08, 2.2, t["accent"])
        self._text(s, MX + 0.35, 1.6, W - 2 * MX - 0.35, 2.2, d.get("statement") or d.get("lead", ""), 24,
                   t["on_dark"], heading=True, anchor=MSO_ANCHOR.MIDDLE, min_pt=14)
        return s

    def bullets(self, d: dict[str, Any]):
        s, t = self._slide(dark=False), self.t
        y0 = self._header(s, d.get("kicker", ""), d["title"], d.get("lead"), dark=False)
        items = d.get("bullets") or []
        avail = 4.95 - y0 - 0.1
        row = min(0.62, avail / max(1, len(items)))
        size = min(fit_size(b, W - 2 * MX - 0.9, row - 0.08, 16, 10) for b in items) if items else 14
        for i, b in enumerate(items):
            y = y0 + 0.1 + i * row
            self._box(s, MX + 0.05, y + 0.1, 0.16, 0.16, t["accent"])
            self._text(s, MX + 0.45, y, W - 2 * MX - 0.5, row - 0.06, b, size, t["ink"], fit=False)
        return s

    def cards(self, d: dict[str, Any]):
        s, t = self._slide(dark=False), self.t
        y0 = self._header(s, d.get("kicker", ""), d["title"], d.get("lead"), dark=False)
        items = (d.get("cards") or [])[:4]
        n = max(1, len(items))
        gap = 0.25
        cw = (W - 2 * MX - gap * (n - 1)) / n
        ch = min(2.6, 4.95 - y0 - 0.2)
        for i, c in enumerate(items):
            x = MX + i * (cw + gap)
            y = y0 + 0.15
            self._box(s, x, y, cw, ch, t["card"] if i % 2 == 0 else t["card_soft"])
            self._badge(s, x + 0.28, y + 0.28, i + 1)
            self._text(s, x + 0.28, y + 0.9, cw - 0.5, 0.5, c["title"], 14.5, t["ink"], heading=True, min_pt=11)
            self._text(s, x + 0.28, y + 1.42, cw - 0.5, ch - 1.6, c.get("text", ""), 11.5, t["body"], min_pt=9)
        return s

    def stats(self, d: dict[str, Any]):
        s, t = self._slide(dark=False), self.t
        y0 = self._header(s, d.get("kicker", ""), d["title"], d.get("lead"), dark=False)
        items = (d.get("stats") or [])[:3]
        n = max(1, len(items))
        gap = 0.3
        cw = (W - 2 * MX - gap * (n - 1)) / n
        y = y0 + 0.2
        for i, st in enumerate(items):
            last = i == n - 1 and n > 1
            x = MX + i * (cw + gap)
            self._box(s, x, y, cw, 1.45, t["accent"] if last else t["card"])
            self._text(s, x + 0.25, y + 0.2, cw - 0.5, 0.6, st["value"], 30,
                       "FFFFFF" if last else t["accent"], heading=True, min_pt=18)
            self._text(s, x + 0.25, y + 0.85, cw - 0.5, 0.5, st.get("label", ""), 11,
                       t["on_dark_soft"] if last else t["body"], min_pt=9)
        if d.get("note"):
            self._box(s, MX, y + 1.75, W - 2 * MX, 0.7, t["card_soft"])
            self._text(s, MX + 0.3, y + 1.85, W - 2 * MX - 0.6, 0.5, d["note"], 12, t["ink"], min_pt=9,
                       anchor=MSO_ANCHOR.MIDDLE)
        return s

    def quote(self, d: dict[str, Any]):
        s, t = self._slide(dark=False), self.t
        y0 = self._header(s, d.get("kicker", ""), d["title"], d.get("lead"), dark=False)
        q = d.get("quote") or {}
        self._box(s, MX, y0 + 0.15, W - 2 * MX, 2.2, t["dark_bg"])
        self._box(s, MX, y0 + 0.15, 0.08, 2.2, t["accent"])
        self._text(s, MX + 0.45, y0 + 0.4, W - 2 * MX - 0.8, 1.4, f"«{q.get('text', '')}»", 20,
                   t["on_dark"], heading=True, min_pt=12, anchor=MSO_ANCHOR.MIDDLE)
        if q.get("source"):
            self._text(s, MX + 0.45, y0 + 1.9, W - 2 * MX - 0.8, 0.3, q["source"], 9.5, t["on_dark_soft"],
                       bold=True, caps=True, fit=False)
        if d.get("note"):
            self._text(s, MX, y0 + 2.6, W - 2 * MX, 0.6, d["note"], 12, t["body"], min_pt=9)
        return s

    def comparison(self, d: dict[str, Any]):
        s, t = self._slide(dark=False), self.t
        y0 = self._header(s, d.get("kicker", ""), d["title"], d.get("lead"), dark=False)
        cols = [d.get("left") or {}, d.get("right") or {}]
        gap = 0.3
        cw = (W - 2 * MX - gap) / 2
        ch = 4.95 - y0 - 0.2
        for i, col in enumerate(cols):
            x = MX + i * (cw + gap)
            y = y0 + 0.15
            dark = i == 1
            self._box(s, x, y, cw, ch, t["card_dark"] if dark else t["card"])
            self._text(s, x + 0.3, y + 0.25, cw - 0.6, 0.35, col.get("title", ""), 14.5,
                       t["on_dark"] if dark else t["ink"], heading=True, min_pt=11)
            items = col.get("items") or []
            row = min(0.55, (ch - 0.9) / max(1, len(items)))
            for k, it in enumerate(items):
                self._text(s, x + 0.3, y + 0.8 + k * row, cw - 0.6, row - 0.05, f"— {it}", 12,
                           t["on_dark_soft"] if dark else t["body"], min_pt=9)
        return s

    def closing(self, d: dict[str, Any]):
        s, t = self._slide(dark=True), self.t
        self._header(s, d.get("kicker", ""), d["title"], None, dark=True)
        items = (d.get("bullets") or [])[:4]
        row = min(1.0, 3.4 / max(1, len(items)))
        for i, b in enumerate(items):
            y = 1.6 + i * row
            self._badge(s, MX + 0.26, y, i + 1)
            self._text(s, MX + 0.88, y, 3.9, row - 0.1, b, 13, t["on_dark"], min_pt=9)
        if d.get("statement"):
            self._text(s, 5.55, 1.5, 3.85, 0.6, d["statement"], 24, t["on_dark"], heading=True, min_pt=14)
        if d.get("note"):
            self._box(s, 5.55, 3.2, 3.85, 1.6, t["accent"])
            self._text(s, 5.81, 3.4, 3.33, 1.2, d["note"], 14, t["on_dark"], heading=True, min_pt=10)
        return s

    # ---------- збирання ----------

    LAYOUTS = ("title", "statement", "bullets", "cards", "stats", "quote", "comparison", "closing")
    DARK = {"title", "statement", "closing"}

    def add(self, d: dict[str, Any], n: int):
        s = getattr(self, d["layout"])(d)
        self._footer(s, n, d["layout"] in self.DARK)
        self._notes(s, d["narration"])

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        return path


def render_deck(plan: dict[str, Any], out: Path, theme: str = "forest") -> Path:
    deck = Deck(theme, plan.get("footer") or plan.get("title", ""))
    for i, slide in enumerate(plan["slides"], start=1):
        deck.add(slide, i)
    return deck.save(out)
