"""Генерує тестову презентацію з «пастками».

Слайди: титульний (темний фон), таблиця, діаграма, прихований, пункти.
Видимих — 4, усього — 5.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

FIXTURE_DIR = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "generated"


def build(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    title_only = prs.slide_layouts[5]

    # 1. Титульний, темний фон
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(0x0B, 0x1F, 0x3A)
    tb = s.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11.3), Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    p.text = "Тестова презентація АБК"
    p.font.size = Pt(48)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 2. Таблиця
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Результати за сценаріями"
    rows, cols = 4, 3
    tbl = s.shapes.add_table(rows, cols, Inches(1), Inches(1.8), Inches(11.3), Inches(3.5)).table
    data = [
        ["Сценарій", "Виручка, млн", "Маржа, %"],
        ["Базовий", "120", "18"],
        ["Оптимістичний", "145", "22"],
        ["Песимістичний", "98", "12"],
    ]
    for r in range(rows):
        for c in range(cols):
            tbl.cell(r, c).text = data[r][c]

    # 3. Діаграма
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Динаміка виручки"
    cd = CategoryChartData()
    cd.categories = ["2023", "2024", "2025", "2026"]
    cd.add_series("Виручка", (98, 110, 120, 131))
    s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.8), Inches(11.3), Inches(5), cd
    )

    # 4. Прихований
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "ПРИХОВАНИЙ СЛАЙД — не має потрапити у відео"
    s._element.set("show", "0")

    # 5. Пункти
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Висновки"
    body = s.placeholders[1].text_frame
    body.text = "Базовий сценарій стабільний"
    for line in ["Оптимістичний потребує інвестицій", "Песимістичний — план ризиків"]:
        body.add_paragraph().text = line

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


# Озвучка видимих слайдів (1, 2, 3, 5). «Наступний слайд» — вербальний маркер переходу.
NARRATION = [
    "Вітаю! Сьогодні я коротко розповім про результати компанії АБК за трьома сценаріями розвитку.",
    "Наступний слайд. Тут видно таблицю результатів. Базовий сценарій дає виручку сто двадцять "
    "мільйонів і маржу вісімнадцять відсотків. Оптимістичний сценарій — сто сорок п'ять мільйонів. "
    "Песимістичний — дев'яносто вісім мільйонів.",
    "Наступний слайд. На діаграмі — динаміка виручки з двох тисяч двадцять третього по двох тисяч "
    "двадцять шостий рік. Виручка стабільно зростає щороку.",
    "Наступний слайд. Підсумуємо. Базовий сценарій стабільний, оптимістичний потребує інвестицій, "
    "а для песимістичного ми готуємо план ризиків. Дякую за увагу.",
]
VISIBLE_SLIDES = [1, 2, 3, 5]
PAUSE_S = 1.2
AUDIO_OFFSET_S = 0.25  # аудіопотік стартує пізніше за відео (пастка start_time)


def build_presenter(out_dir: Path) -> dict:
    """Синтетичне HDR (HLG) відео ведучого з українською озвучкою + еталонні таймкоди слайдів."""
    import json
    import subprocess

    out_dir.mkdir(parents=True, exist_ok=True)
    tmp = out_dir.parent / "media_tmp"  # проміжні файли поза каталогом входу
    tmp.mkdir(parents=True, exist_ok=True)
    parts, starts, t = [], [], 0.0
    for i, text in enumerate(NARRATION):
        aiff = tmp / f"seg{i}.aiff"
        subprocess.run(["say", "-v", "Lesya", "-o", str(aiff), text], check=True)
        wav = tmp / f"seg{i}.wav"
        subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-i", str(aiff),
                        "-af", f"apad=pad_dur={PAUSE_S}", "-ar", "48000", "-ac", "1", str(wav)], check=True)
        dur = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration",
                                    "-of", "csv=p=0", str(wav)], capture_output=True, text=True,
                                   check=True).stdout)
        starts.append(round(t, 3))
        t += dur
        parts.append(wav)
    listing = tmp / "list.txt"
    listing.write_text("".join(f"file '{p.name}'\n" for p in parts))
    narration = tmp / "narration.wav"
    subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-f", "concat", "-safe", "0",
                    "-i", str(listing), "-c", "copy", str(narration)], check=True)

    video = out_dir / "presenter.mov"
    subprocess.run([
        "ffmpeg", "-hide_banner", "-loglevel", "error", "-y",
        "-f", "lavfi", "-i", f"testsrc2=size=1280x720:rate=30:duration={t:.3f}",
        "-itsoffset", str(AUDIO_OFFSET_S), "-i", str(narration),
        "-map", "0:v", "-map", "1:a",
        "-c:v", "libx265", "-pix_fmt", "yuv420p10le",
        "-x265-params", "colorprim=bt2020:transfer=arib-std-b67:colormatrix=bt2020nc:log-level=error",
        "-tag:v", "hvc1", "-color_primaries", "bt2020", "-color_trc", "arib-std-b67",
        "-colorspace", "bt2020nc", "-c:a", "aac", "-b:a", "160k", str(video),
    ], check=True)
    truth = {
        "audio_offset_s": AUDIO_OFFSET_S,
        "duration_s": round(t + AUDIO_OFFSET_S, 3),
        # момент початку фрази для кожного видимого слайда, у часі відео
        "transitions": [{"slide": s, "start": round(st + AUDIO_OFFSET_S, 3)}
                        for s, st in zip(VISIBLE_SLIDES, starts)],
    }
    (out_dir / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2))
    return truth


SCRIPT_MD = """# Сценарій тестової лекції
Текст до першого заголовка ігнорується.

## Слайд 1
Вітаю! Сьогодні коротко про результати компанії АБК за трьома сценаріями розвитку.

## Слайд 2
Тут таблиця результатів. [пауза]
Базовий сценарій дає виручку сто двадцять мільйонів і маржу вісімнадцять відсотків.
Оптимістичний — сто сорок п'ять мільйонів, песимістичний — дев'яносто вісім.

## Слайд 3
На діаграмі — динаміка виручки з двох тисяч двадцять третього по двох тисяч двадцять шостий рік.

## Слайд 5
Підсумуємо. Базовий сценарій стабільний, оптимістичний потребує інвестицій,
а для песимістичного ми готуємо план ризиків. [пауза 1.5] Дякую за увагу!
"""


def build_script_fixture() -> Path:
    target = FIXTURE_DIR / "script_input"
    build(target / "presentation.pptx")
    (target / "script.md").write_text(SCRIPT_MD, encoding="utf-8")
    return target


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--script":
        print(build_script_fixture())
        raise SystemExit
    if len(sys.argv) > 1 and sys.argv[1] == "--media":
        target = FIXTURE_DIR / "job_input"
        build(target / "presentation.pptx")
        print(build_presenter(target))
    else:
        out = Path(sys.argv[1]) if len(sys.argv) > 1 else FIXTURE_DIR / "phase0_deck.pptx"
        print(build(out))
