"""slide_manifest.json з PPTX: елементи з element_id і bbox, таблиці, діаграми, нотатки, анімації."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
PRESET_CLASS = {"entr": "entrance", "exit": "exit", "emph": "emphasis", "path": "motion"}


class _Xform:
    """Перетворення координат групи (child space → slide EMU → px)."""

    def __init__(self, scale: float, ox: float = 0, oy: float = 0, sx: float = 1, sy: float = 1,
                 cx: float = 0, cy: float = 0):
        self.scale, self.ox, self.oy, self.sx, self.sy, self.cx, self.cy = scale, ox, oy, sx, sy, cx, cy

    def bbox(self, left, top, width, height) -> list[float] | None:
        if None in (left, top, width, height):
            return None
        x = self.ox + (left - self.cx) * self.sx
        y = self.oy + (top - self.cy) * self.sy
        return [round(x * self.scale, 1), round(y * self.scale, 1),
                round(width * self.sx * self.scale, 1), round(height * self.sy * self.scale, 1)]

    def child(self, group) -> "_Xform":
        xfrm = group._element.find(".//p:grpSpPr/a:xfrm", NS)
        if xfrm is None:
            return self
        off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
        choff, chext = xfrm.find("a:chOff", NS), xfrm.find("a:chExt", NS)
        if None in (off, ext, choff, chext):
            return self
        ox = self.ox + (int(off.get("x")) - self.cx) * self.sx
        oy = self.oy + (int(off.get("y")) - self.cy) * self.sy
        sx = self.sx * (int(ext.get("cx")) / max(1, int(chext.get("cx"))))
        sy = self.sy * (int(ext.get("cy")) / max(1, int(chext.get("cy"))))
        return _Xform(self.scale, ox, oy, sx, sy, int(choff.get("x")), int(choff.get("y")))


def _text(shape) -> str:
    return shape.text_frame.text.strip() if getattr(shape, "has_text_frame", False) and shape.has_text_frame else ""


def _shape_elements(shape, slide_no: int, xf: _Xform) -> list[dict[str, Any]]:
    sid = f"slide{slide_no}/shape{shape.shape_id}"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        out = []
        cxf = xf.child(shape)
        for sub in shape.shapes:
            for el in _shape_elements(sub, slide_no, cxf):
                el.setdefault("parent", sid)
                out.append(el)
        return out

    bbox = xf.bbox(shape.left, shape.top, shape.width, shape.height)
    el: dict[str, Any] = {"element_id": sid, "type": "shape", "bbox": bbox, "text": None, "parent": None}

    if getattr(shape, "has_table", False) and shape.has_table:
        tbl = shape.table
        cells = [[tbl.cell(r, c).text.strip() for c in range(len(tbl.columns))] for r in range(len(tbl.rows))]
        el.update(type="table", table={"rows": len(tbl.rows), "cols": len(tbl.columns), "cells": cells})
        out = [el]
        y = shape.top
        for r, row in enumerate(tbl.rows):
            x = shape.left
            for c, col in enumerate(tbl.columns):
                out.append({
                    "element_id": f"{sid}/r{r}c{c}", "type": "table_cell", "parent": sid,
                    "bbox": xf.bbox(x, y, col.width, row.height), "text": cells[r][c],
                })
                x += col.width
            y += row.height
        return out

    if getattr(shape, "has_chart", False) and shape.has_chart:
        chart = shape.chart
        plot = chart.plots[0] if len(chart.plots) else None
        el.update(type="chart", text=chart.chart_title.text_frame.text if chart.has_title else None, chart={
            "chart_type": str(chart.chart_type).split(".")[-1].split(" ")[0],
            "categories": [str(c) for c in plot.categories] if plot else [],
            "series": [{"name": s.name, "values": [None if v is None else float(v) for v in s.values]}
                       for s in (plot.series if plot else [])],
        })
        return [el]

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        el["type"] = "picture"
        return [el]

    text = _text(shape)
    if shape.is_placeholder and shape.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
        el["type"] = "title"
    elif text:
        el["type"] = "text"
    elif not text:
        return []  # декоративні фігури без тексту не є цілями режисури
    el["text"] = text
    out = [el]
    if el["type"] == "text":
        for i, p in enumerate(shape.text_frame.paragraphs):
            ptext = "".join(r.text for r in p.runs).strip()
            if ptext:
                # bbox абзаців — Фаза 2 (pdftotext -bbox-layout)
                out.append({"element_id": f"{sid}/p{i}", "type": "paragraph", "parent": sid,
                            "bbox": None, "text": ptext})
    return out


def _builds(slide, slide_no: int) -> list[dict[str, Any]]:
    timing = slide._element.find("p:timing", NS)
    if timing is None:
        return []
    builds, seen = [], set()
    for tgt in timing.iter(f"{{{NS['p']}}}spTgt"):
        spid = tgt.get("spid")
        para = tgt.find("p:txEl/p:pRg", NS)
        key = (spid, para.get("st") if para is not None else None)
        if key in seen:
            continue
        seen.add(key)
        kind = "entrance"
        for anc in tgt.iterancestors(f"{{{NS['p']}}}cTn"):
            if anc.get("presetClass"):
                kind = PRESET_CLASS.get(anc.get("presetClass"), "entrance")
                break
        eid = f"slide{slide_no}/shape{spid}" + (f"/p{key[1]}" if key[1] is not None else "")
        builds.append({"order": len(builds) + 1, "element_id": eid, "kind": kind,
                       "level": "paragraph" if para is not None else "shape"})
    return builds


def _advance_after(slide) -> float | None:
    for el in slide._element.iter():
        if isinstance(el.tag, str) and el.get("advTm"):
            return int(el.get("advTm")) / 1000
    return None


def build_manifest(pptx: Path, png_width: int, png_height: int) -> dict[str, Any]:
    prs = Presentation(str(pptx))
    scale = png_width / Emu(prs.slide_width)
    xf = _Xform(scale)
    slides = []
    for n, slide in enumerate(prs.slides, start=1):
        elements: list[dict[str, Any]] = []
        for shape in slide.shapes:
            elements.extend(_shape_elements(shape, n, xf))
        notes = None
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip() or None
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else None
        if not title:
            title = next((e["text"] for e in elements if e["type"] in ("title", "text") and e["text"]), None)
        slides.append({
            "slide": n,
            "hidden": slide._element.get("show") == "0",
            "pdf_page": None,
            "image": None,
            "content_mask": None,
            "title": title,
            "notes": notes,
            "elements": elements,
            "builds": _builds(slide, n),
            "advance_after_s": _advance_after(slide),
        })
    return {"source": str(pptx), "renderer": "powerpoint", "slide_width_px": png_width,
            "slide_height_px": png_height, "slides": slides}
