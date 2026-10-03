"""Режим «текст → презентація»: читання тексту, перевірка плану, збирання PPTX (без Claude)."""

import zipfile

import pytest
from pptx import Presentation

from scripts import state
from scripts.compose import check_plan, read_lecture, target_slides
from scripts.deck_render import fit_size, render_deck
from scripts.state import StepError, discover_inputs

NARR = ("Сьогодні говоримо про те, як писати запити до моделей так, щоб отримувати потрібний результат "
        "з першого або другого разу. Розберемо формулу з чотирьох частин і кілька принципів, "
        "які однаково радять усі великі розробники. ") * 2


def _slide(layout, **kw):
    base = {"layout": layout, "kicker": "Розділ", "title": "Заголовок", "lead": None, "meta": None,
            "statement": None, "note": None, "bullets": [], "cards": [], "stats": [], "quote": None,
            "left": None, "right": None, "narration": NARR}
    base.update(kw)
    return base


def _plan():
    return {"title": "Лекція", "footer": "Курс · Лекція", "slides": [
        _slide("title", lead="Підзаголовок"),
        _slide("bullets", bullets=["Перший пункт", "Другий пункт", "Третій пункт"]),
        _slide("cards", cards=[{"title": "Ціль", "text": "Що має вийти."}, {"title": "Контекст", "text": "Для кого."}]),
        _slide("stats", stats=[{"value": "4", "label": "вендори"}, {"value": "60+", "label": "сторінок"}]),
        _slide("quote", quote={"text": "Пишіть як для нового колеги", "source": "Посібник"}),
        _slide("comparison", left={"title": "До", "items": ["а", "б"]}, right={"title": "Після", "items": ["в", "г"]}),
        _slide("statement", statement="Головна думка."),
        _slide("closing", bullets=["Висновок один", "Висновок два"], statement="Дякую."),
    ]}


def test_valid_plan_passes():
    words = sum(len(s["narration"].split()) for s in _plan()["slides"])
    assert check_plan(_plan(), words, (6, 10)) == []


def test_plan_errors_are_specific():
    plan = _plan()
    plan["slides"][0]["layout"] = "bullets"
    plan["slides"][1]["bullets"] = ["один"]
    plan["slides"][2]["narration"] = "Виручка зросла на 20 відсотків."
    plan["slides"][3]["stats"] = [{"value": "1234567890", "label": "x"}]
    errors = " | ".join(check_plan(plan, 300, (6, 10)))
    assert "slide 1 must use layout 'title'" in errors
    assert "needs 3–6 bullets" in errors
    assert "digits" in errors and "narration has" in errors
    assert "stat value" in errors


def test_plan_that_drops_content_is_rejected():
    errors = check_plan(_plan(), 5000, (6, 10))
    assert any("cut too much" in e for e in errors)


def test_target_slides_scales_with_text():
    assert target_slides(130 * 20, {}) == (18, 23)
    assert target_slides(100, {})[0] == 3


def test_read_lecture_md_and_docx(tmp_path):
    md = tmp_path / "lecture.md"
    md.write_text("﻿" + "Слово " * 100, encoding="utf-8")
    assert read_lecture(md).startswith("Слово")
    docx = tmp_path / "lecture.docx"
    paras = "".join(f"<w:p><w:r><w:t>Абзац {i} {'текст ' * 20}</w:t></w:r></w:p>" for i in range(5))
    with zipfile.ZipFile(docx, "w") as z:
        z.writestr("word/document.xml", f"<w:document><w:body>{paras}</w:body></w:document>")
    text = read_lecture(docx)
    assert "Абзац 0" in text and "Абзац 4" in text
    short = tmp_path / "short.txt"
    short.write_text("Замало.", encoding="utf-8")
    with pytest.raises(StepError, match="замало"):
        read_lecture(short)


def test_discover_lecture_only(tmp_path):
    (tmp_path / "Лекція про промптинг.docx").write_bytes(b"x")
    found = discover_inputs(tmp_path)
    assert set(found) == {"lecture"} and state.job_mode(found) == "script"


def test_discover_presentation_with_docx_keeps_notes_mode(tmp_path):
    (tmp_path / "deck.pptx").write_bytes(b"x")
    (tmp_path / "сценарій.docx").write_bytes(b"x")
    found = discover_inputs(tmp_path)
    assert set(found) == {"presentation"}  # docx поруч із презентацією не чіпаємо


def test_render_deck_has_all_slides_and_notes(tmp_path):
    out = render_deck(_plan(), tmp_path / "deck.pptx")
    prs = Presentation(str(out))
    assert len(prs.slides) == 8
    assert all(s.notes_slide.notes_text_frame.text == NARR for s in prs.slides)
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Курс · Лекція" in texts and "Пишіть як для нового колеги" in texts


def test_fit_size_shrinks_long_text():
    assert fit_size("Коротко", 8, 0.5, 27) == 27
    assert fit_size("Дуже довгий заголовок " * 6, 8, 0.5, 27) < 27


@pytest.mark.parametrize("title", ["Коротко", "Промптинг 2026: як писати запити до AI",
                                   "Дуже довгий заголовок лекції про те, як писати запити до моделей штучного "
                                   "інтелекту у двадцять шостому році"])
def test_title_slide_lead_never_overlaps_title(tmp_path, title):
    plan = {"title": "x", "footer": "", "slides": [_slide("title", title=title, lead="Підзаголовок лекції",
                                                          meta="Опис")]}
    prs = Presentation(str(render_deck(plan, tmp_path / "t.pptx")))
    boxes = {sh.text_frame.text: sh for sh in prs.slides[0].shapes if sh.has_text_frame}
    t, lead, meta = boxes[title], boxes["Підзаголовок лекції"], boxes["Опис"]
    assert lead.top >= t.top + t.height
    assert meta.top >= lead.top + lead.height
